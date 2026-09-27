"""python-pptx implementation of PowerPoint parsing."""

from io import BytesIO

from pptx import Presentation

from rag.parsers.powerpoint.powerpoint_parser import PowerPointParser


class PythonPptxPowerPointParser(PowerPointParser):
    """Extract slide text with python-pptx."""

    def parse(self, content: bytes) -> str:
        presentation = Presentation(BytesIO(content))
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            text = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if text:
                slides.append(f"# Slide {index}\n" + "\n".join(text))
        return "\n\n".join(slides)
