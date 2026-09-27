"""python-docx implementation of DOCX parsing."""

from io import BytesIO

from docx import Document

from rag.parsers.word.docx_parser import DocxParser


class PythonDocxDocxParser(DocxParser):
    """Extract DOCX paragraphs and tables with python-docx."""

    def parse(self, content: bytes) -> str:
        document = Document(BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
        return "\n".join(parts)
