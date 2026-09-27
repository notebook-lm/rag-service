"""openpyxl implementation of Excel parsing."""

from io import BytesIO

from openpyxl import load_workbook

from rag.parsers.excel.excel_parser import ExcelParser


class OpenPyxlExcelParser(ExcelParser):
    """Extract non-empty workbook rows with openpyxl."""

    def parse(self, content: bytes) -> str:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        parts: list[str] = []
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                cells = [str(value) if value is not None else "" for value in row]
                if any(cell.strip() for cell in cells):
                    rows.append("\t".join(cells).rstrip())
            if rows:
                parts.append(f"# {worksheet.title}\n" + "\n".join(rows))
        return "\n\n".join(parts)
